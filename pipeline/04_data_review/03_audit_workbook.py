"""审查自动标准的Excel。"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment


空值 = {"", "N", "n", "无", "未提及", "[]", "null", "None"}

中药名规范映射 = {
    "麸炒白术": "白术",
    "炒白术": "白术",
    "炒黄芩": "黄芩",
    "制附片": "制附子",
    "熟附子": "附子",
    "鲜生姜": "生姜",
    "炙甘草": "甘草",
    "生甘草": "甘草",
    "生草": "甘草",
    "姜半夏": "半夏",
    "法半夏": "半夏",
    "法夏": "半夏",
    "清半夏": "半夏",
    "制半夏": "半夏",
    "炒半夏": "半夏",
    "炒薏苡仁": "薏苡仁",
    "炒枣仁": "酸枣仁",
    "炒山药": "山药",
    "全蝎粉": "全蝎",
    "白芍药": "白芍",
    "大白芍": "白芍",
    "酒白芍": "白芍",
    "云茯苓": "茯苓",
    "广陈皮": "陈皮",
    "全当归": "当归",
    "酒当归": "当归",
    "煨姜": "生姜",
    "红枣": "大枣",
    "枣": "大枣",
    "炮姜": "干姜",
    "淡干姜": "干姜",
    "千姜": "干姜",
    "于姜": "干姜",
    "生栀子": "栀子",
    "栀仁": "栀子",
    "白蔻仁": "白豆蔻",
    "白寇仁": "白豆蔻",
    "瓜蒌仁": "瓜蒌",
    "生龙骨": "龙骨",
    "元明粉": "芒硝",
    "玄明粉": "芒硝",
    "生晒参": "人参",
    "红参": "人参",
    "银花": "金银花",
    "麦门冬": "麦冬",
    "辽细辛": "细辛",
    "川厚朴": "厚朴",
    "川连": "黄连",
    "连": "黄连",
    "芩": "黄芩",
    "生大黄": "大黄",
    "酒大黄": "大黄",
    "生黄芪": "黄芪",
    "江枳实": "枳实",
    "炒枳壳": "枳壳",
    "炒枳实": "枳实",
    "附片": "附子",
    "肉桂心": "肉桂",
    "乌梅肉": "乌梅",
    "川棟子": "川楝子",
    "炙草": "甘草",
    "炒莱菔子": "莱菔子",
    "炙枇杷叶": "枇杷叶",
    "煅瓦楞": "瓦楞子",
    "生石膏": "石膏",
    "鲜苦楝根皮": "苦楝根皮",
    "吴萸": "吴茱萸",
    "广木香": "木香",
    "良姜": "高良姜",
    "元胡": "延胡索",
    "乌贼骨": "海螵蛸",
    "浙贝": "浙贝母",
    "丹皮": "牡丹皮",
    "苏叶": "紫苏叶",
    "萎蕤": "玉竹",
    "芍药": "白芍",
}

表达级中药规范 = {
    "制附子钱半": "附子",
    "制附子錢半": "附子",
    "熟附子钱半": "附子",
    "熟附子錢半": "附子",
}

病名替换 = {
    "口疮之病": "口疮",
    "吐狂血": "吐血",
}
移出病名 = {"渗流中道", "渗留中道"}

剂量单位模式 = re.compile(
    r"^(?P<name>.+?)(?P<dose_unit>(?:\d+(?:\.\d+)?|[零一二三四五六七八九十百千万半两]+)"
    r"(?:mg|mL|ml|g|克|毫克|枚|片|个|条|两|钱|錢|分|厘|升)"
    r"(?:[零一二三四五六七八九十百千万半两]+(?:钱|錢|分|厘))?)(?P<suffix>.*)$"
)
钱半模式 = re.compile(r"^(?P<name>.+?)(?:钱半|錢半)(?P<suffix>.*)$")
裸单位尾缀模式 = re.compile(r"^(?P<name>.+?)(?:mg|mL|ml|g|克|毫克)$")

重点异常 = [
    "肉桂心",
    "半夏钱半",
    "制附子钱半",
    "附子钱半",
    "人参钱半",
    "干姜七分",
    "干姜一錢",
    "口疮之病",
    "吐狂血",
    "渗流中道",
    "渗留中道",
]

备注替换 = {
    "方剂组成：熟附子->制附子": "方剂组成：熟附子->附子（钱半为剂量表达，不进入中药名）",
}


def 清理值(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text in 空值 else text


def 清理术语(value: Any) -> str:
    text = 清理值(value).replace("\u3000", " ")
    text = re.sub(r"\s+", "", text)
    return text.strip("“”‘’\"'，,。；;：: ")


def 拆分多值(text: str) -> list[str]:
    if not 清理值(text):
        return []
    parts = re.split(r"[、；;\n/]+", str(text))
    result: list[str] = []
    seen: set[str] = set()
    for part in parts:
        term = 清理术语(part)
        if term and term not in seen:
            result.append(term)
            seen.add(term)
    return result


def 合并多值(parts: list[str]) -> str:
    return "、".join(parts) if parts else "N"


def 追加备注(old: str, new_note: str) -> str:
    old = 清理值(old)
    if not old:
        return new_note
    if new_note in old:
        return old
    return f"{old}；{new_note}"


def 清理备注(old: str) -> str:
    text = 清理值(old)
    for src, dst in 备注替换.items():
        text = text.replace(src, dst)
    return text


def 规范中药名(name: str) -> str:
    clean = 清理术语(name)
    if not clean:
        return ""
    clean = re.sub(r"生地(?!黄)", "生地黄", clean)
    clean = re.sub(r"熟地(?!黄)", "熟地黄", clean)
    return 中药名规范映射.get(clean, clean)


def 提取中药名(expression: str) -> tuple[str, str, bool]:
    text = 清理术语(expression)
    if not text:
        return "", "", False

    if text in 表达级中药规范:
        return text, 表达级中药规范[text], True

    match = 钱半模式.match(text) or 剂量单位模式.match(text) or 裸单位尾缀模式.match(text)
    if match:
        original_name = 清理术语(match.group("name"))
        return original_name, 规范中药名(original_name), True

    return text, 规范中药名(text), False


def 规范方剂组成(value: str) -> tuple[str, list[str]]:
    terms = 拆分多值(value)
    if not terms:
        return 清理值(value), []

    normalized_terms: list[str] = []
    seen: set[str] = set()
    changed_terms: list[str] = []
    removed_dose = False
    source_name_count = 0

    for term in terms:
        original_name, normalized_name, had_dose = 提取中药名(term)
        if had_dose:
            removed_dose = True
        if not normalized_name:
            continue
        source_name_count += 1
        if normalized_name not in seen:
            normalized_terms.append(normalized_name)
            seen.add(normalized_name)
        if term != normalized_name:
            if original_name and original_name != normalized_name:
                changed_terms.append(f"{original_name}->{normalized_name}")
            elif had_dose:
                changed_terms.append(f"{term}->{normalized_name}")

    dedup_changed_terms = list(dict.fromkeys(changed_terms))
    notes: list[str] = []
    if removed_dose:
        notes.append("去除剂量表达，仅保留规范中药名")
    if len(normalized_terms) < source_name_count:
        notes.append("同名中药去重")
    notes.extend(dedup_changed_terms)
    return 合并多值(normalized_terms), notes


def 规范病名(value: str) -> tuple[str, list[str]]:
    terms = 拆分多值(value)
    if not terms:
        return 清理值(value), []

    changed: list[str] = []
    result: list[str] = []
    seen: set[str] = set()
    for term in terms:
        if term in 移出病名:
            changed.append(f"{term}->N")
            continue
        normalized = 病名替换.get(term, term)
        if normalized != term:
            changed.append(f"{term}->{normalized}")
        if normalized and normalized not in seen:
            result.append(normalized)
            seen.add(normalized)
    return 合并多值(result), changed


def find_header(headers: dict[str, int], candidates: list[str]) -> int:
    for name in candidates:
        if name in headers:
            return headers[name]
    raise KeyError(f"缺少列：{candidates}")


def 审查或修正(输入路径: Path, apply: bool, 输出路径: Path | None) -> dict[str, Any]:
    workbook = load_workbook(输入路径)
    sheet = workbook.active
    headers = {
        str(cell.value).strip(): index + 1
        for index, cell in enumerate(sheet[1])
        if cell.value is not None and str(cell.value).strip()
    }
    case_col = find_header(headers, ["医案编号"])
    formula_col = find_header(headers, ["方剂组成"])
    tcm_disease_col = find_header(headers, ["中医病名"])
    remark_col = find_header(headers, ["标准化备注", "备注"])

    rows_report: list[dict[str, Any]] = []
    focus_hits: list[dict[str, str]] = []

    for row_index in range(2, sheet.max_row + 1):
        case_id = 清理值(sheet.cell(row_index, case_col).value)
        if not case_id:
            continue

        formula_old = 清理值(sheet.cell(row_index, formula_col).value)
        disease_old = 清理值(sheet.cell(row_index, tcm_disease_col).value)
        remark_old = 清理值(sheet.cell(row_index, remark_col).value)
        remark_clean = 清理备注(remark_old)

        formula_new, formula_notes = 规范方剂组成(formula_old)
        disease_new, disease_notes = 规范病名(disease_old)

        note_parts: list[str] = []
        if formula_new != formula_old and formula_notes:
            note_parts.append("方剂组成：" + "，".join(formula_notes))
        if disease_new != disease_old and disease_notes:
            note_parts.append("中医病名：" + "，".join(disease_notes))

        if note_parts or remark_clean != remark_old:
            rows_report.append({
                "row": row_index,
                "case": case_id,
                "方剂组成_before": formula_old,
                "方剂组成_after": formula_new,
                "中医病名_before": disease_old,
                "中医病名_after": disease_new,
                "备注_before": remark_old if remark_clean != remark_old else "",
                "备注_after": remark_clean if remark_clean != remark_old else "",
                "备注追加": "；".join(note_parts),
            })
            if apply:
                sheet.cell(row_index, formula_col).value = formula_new
                sheet.cell(row_index, tcm_disease_col).value = disease_new
                next_remark = remark_clean
                if note_parts:
                    next_remark = 追加备注(next_remark, "；".join(note_parts))
                sheet.cell(row_index, remark_col).value = next_remark

        joined_after = "\n".join([formula_new, disease_new])
        for term in 重点异常:
            if term in joined_after:
                focus_hits.append({
                    "case": case_id,
                    "term": term,
                    "方剂组成": formula_new,
                    "中医病名": disease_new,
                })

    if apply:
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        assert 输出路径 is not None
        输出路径.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(输出路径)

    return {
        "path": str(输入路径),
        "output": str(输出路径) if apply else "",
        "apply": apply,
        "changes": rows_report,
        "focus_hits_after_planned_change": focus_hits,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="审查标准化表的药名和病名，修订结果需人工确认")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--apply", action="store_true", help="将建议修订另存为新工作簿")
    parser.add_argument("--output", type=Path, help="仅与 --apply 同用；必须为不存在的新文件")
    args = parser.parse_args()
    if args.apply and (args.output is None or args.output.resolve() == args.input.resolve() or args.output.exists()):
        parser.error("--apply 必须配合未存在且不同于输入的 --output")
    if args.output is not None and not args.apply:
        parser.error("--output 仅与 --apply 同用")
    report = 审查或修正(args.input, args.apply, args.output)
    # 默认只报告 Case ID 和修改字段，不把逐案内容打印到日志。
    report["changes"] = [
        {"case": row["case"], "fields": [field for field in ("方剂组成", "中医病名", "备注") if row.get(f"{field}_before") != row.get(f"{field}_after")]}
        for row in report["changes"]
    ]
    report["focus_hits_after_planned_change"] = len(report["focus_hits_after_planned_change"])
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
