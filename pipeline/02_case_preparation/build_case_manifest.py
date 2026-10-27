"""核对人工选定的 Case 文本与切分预览表，并建立可追溯的来源映射。

输入为阶段01经人工筛选的 Excel 和阶段02的 Case*.txt。序号取自表中“序号”，
不按文件排序重新编号；文本不一致时停止，避免把人工替换内容归到错误文档。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook


项目根目录 = Path(__file__).resolve().parents[2]
默认工作簿 = 项目根目录 / "pipeline" / "01_document_processing" / "output" / "全部医案_按分隔线拆分.xlsx"
默认拆分JSON = 默认工作簿.with_suffix(".json")
默认医案目录 = Path(__file__).resolve().parent / "output"
默认来源目录 = 项目根目录 / "data" / "raw"
默认输出路径 = 默认医案目录 / "医案映射清单.csv"
字段 = ["医案编号", "全局序号", "文档名称", "文档内序号", "医案文件", "字符数", "文本SHA256"]


def 标准化比较文本(value: object) -> str:
    """只忽略空白差异；不改写人工选定的临床文本。"""
    return re.sub(r"\s+", "", str(value or ""))


def 解析医案编号(value: object) -> tuple[str, str]:
    text = str(value or "").strip()
    if re.fullmatch(r"\d+(?:\.0)?", text):
        number = int(float(text))
    elif re.fullmatch(r"Case\d+", text):
        number = int(text[4:])
    else:
        raise ValueError("Excel 序号必须是正整数或 Case 编号")
    if number <= 0:
        raise ValueError("Excel 序号必须大于零")
    return f"Case{number:03d}", str(number)


def 读取旧拆分序号(path: Path) -> dict[tuple[str, str], str]:
    """仅对文本仍与拆分 JSON 完全对应的记录恢复原文档内序号。"""
    if not path.exists():
        return {}
    records = json.loads(path.read_text(encoding="utf-8-sig"))
    result: dict[tuple[str, str], str] = {}
    for record in records:
        key = (str(record["文档名称"]), 标准化比较文本(record["具体医案"]))
        if key in result:
            result[key] = "N"  # 同一文档中出现重复文本时不猜测其序号。
        else:
            result[key] = str(record["文档内序号"])
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="核对人工选案文本并生成来源映射")
    parser.add_argument("--workbook", type=Path, default=默认工作簿)
    parser.add_argument("--case-dir", type=Path, default=默认医案目录)
    parser.add_argument("--source-dir", type=Path, default=默认来源目录)
    parser.add_argument("--split-json", type=Path, default=默认拆分JSON)
    parser.add_argument("--output", type=Path, default=默认输出路径)
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError(f"映射文件已存在，请使用新批次输出路径：{args.output}")
    if not args.workbook.is_file() or not args.case_dir.is_dir():
        raise FileNotFoundError("缺少人工选案 Excel 或 Case 文本目录")

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        required = ["序号", "文档名称", "具体医案"]
        if any(name not in headers for name in required):
            raise ValueError(f"选案表必须包含：{', '.join(required)}")
        indexes = {name: headers.index(name) for name in required}
        old_ordinals = 读取旧拆分序号(args.split_json)
        mapping: list[dict[str, str]] = []
        seen: set[str] = set()
        mismatched: list[str] = []
        for row in rows:
            if not any(value is not None and str(value).strip() for value in row):
                continue
            case_id, number = 解析医案编号(row[indexes["序号"]])
            if case_id in seen:
                raise ValueError(f"选案表存在重复编号：{case_id}")
            seen.add(case_id)
            source_name = str(row[indexes["文档名称"]] or "").strip()
            if not source_name or Path(source_name).name != source_name:
                raise ValueError(f"{case_id} 的来源文档名称无效")
            if not (args.source_dir / source_name).is_file():
                raise FileNotFoundError(f"{case_id} 的来源文档不存在：{source_name}")
            case_path = args.case_dir / f"{case_id}.txt"
            if not case_path.is_file():
                raise FileNotFoundError(f"选案表中的医案文本不存在：{case_id}")
            case_text = case_path.read_text(encoding="utf-8-sig")
            excel_text = row[indexes["具体医案"]]
            if not 标准化比较文本(excel_text) or not 标准化比较文本(case_text):
                raise ValueError(f"{case_id} 的选案文本为空")
            if 标准化比较文本(case_text) != 标准化比较文本(excel_text):
                mismatched.append(case_id)
            try:
                case_file = str(case_path.resolve().relative_to(项目根目录))
            except ValueError:
                case_file = str(case_path.resolve())
            mapping.append({
                "医案编号": case_id,
                "全局序号": number,
                "文档名称": source_name,
                "文档内序号": old_ordinals.get((source_name, 标准化比较文本(excel_text)), "N"),
                "医案文件": case_file,
                "字符数": str(len(case_text)),
                "文本SHA256": hashlib.sha256(case_text.encode("utf-8")).hexdigest(),
            })
    finally:
        workbook.close()

    extra = sorted(path.stem for path in args.case_dir.glob("Case*.txt") if path.stem not in seen)
    if extra:
        raise ValueError(f"Case 目录有未列入选案表的文件：{', '.join(extra)}")
    if mismatched:
        raise ValueError(f"人工选案 Excel 与 Case 文本不一致，请核对后再建映射：{', '.join(mismatched)}")
    if not mapping:
        raise ValueError("选案表没有可用记录")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=字段)
        writer.writeheader()
        writer.writerows(mapping)
    print(f"来源映射已生成：{args.output}；记录数：{len(mapping)}")


if __name__ == "__main__":
    main()
