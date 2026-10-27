"""从已人工确认的标准化Excel，按 kg_schema 构建节点、关系和三元组 CSV。

"""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import kg_schema


项目根目录 = Path(__file__).resolve().parents[2]
默认映射路径 = 项目根目录 / "pipeline" / "02_case_preparation" / "output" / "医案映射清单.csv"
默认输出目录 = Path(__file__).resolve().parent / "output"
空值集合 = {"", "N", "n", "无", "未提及", "[]", "null", "None"}

实体字段关系 = [
    ("症状体征", "症状体征", "具有症状体征"),
    ("舌象", "舌象", "具有舌象"),
    ("脉象", "脉象", "具有脉象"),
    ("辅助检查", "辅助检查", "包含辅助检查"),
    ("中医病名", "中医病名", "中医诊断为"),
    ("西医病名", "西医病名", "西医诊断为"),
    ("证型", "证型", "辨证为"),
    ("治则治法", "治则治法", "采用治法"),
]

中药名规范映射 = {
    "麸炒白术": "白术",
    "炒白术": "白术",
    "炒黄芩": "黄芩",
    "制附片": "制附子",
    "熟附子": "附子",
    "鲜生姜": "生姜",
    "炙甘草": "甘草",
    "生甘草": "甘草",
    "生草": "甘草",
    "姜半夏": "半夏",
    "法半夏": "半夏",
    "法夏": "半夏",
    "清半夏": "半夏",
    "制半夏": "半夏",
    "炒半夏": "半夏",
    "炒薏苡仁": "薏苡仁",
    "炒枣仁": "酸枣仁",
    "炒山药": "山药",
    "全蝎粉": "全蝎",
    "白芍药": "白芍",
    "大白芍": "白芍",
    "酒白芍": "白芍",
    "云茯苓": "茯苓",
    "广陈皮": "陈皮",
    "全当归": "当归",
    "酒当归": "当归",
    "煨姜": "生姜",
    "红枣": "大枣",
    "炮姜": "干姜",
    "淡干姜": "干姜",
    "千姜": "干姜",
    "于姜": "干姜",
    "生栀子": "栀子",
    "栀仁": "栀子",
    "白蔻仁": "白豆蔻",
    "白寇仁": "白豆蔻",
    "瓜蒌仁": "瓜蒌",
    "生龙骨": "龙骨",
    "元明粉": "芒硝",
    "玄明粉": "芒硝",
    "生晒参": "人参",
    "红参": "人参",
    "银花": "金银花",
    "麦门冬": "麦冬",
    "辽细辛": "细辛",
    "川厚朴": "厚朴",
    "川连": "黄连",
    "连": "黄连",
    "芩": "黄芩",
    "生大黄": "大黄",
    "酒大黄": "大黄",
    "生黄芪": "黄芪",
    "江枳实": "枳实",
    "炒枳壳": "枳壳",
    "炒枳实": "枳实",
    "附片": "附子",
    "肉桂心": "肉桂",
    "乌梅肉": "乌梅",
    "川棟子": "川楝子",
    "炙草": "甘草",
    "炒莱菔子": "莱菔子",
    "炙枇杷叶": "枇杷叶",
    "煅瓦楞": "瓦楞子",
    "生石膏": "石膏",
    "鲜苦楝根皮": "苦楝根皮",
    "吴萸": "吴茱萸",
    "广木香": "木香",
    "良姜": "高良姜",
    "元胡": "延胡索",
    "乌贼骨": "海螵蛸",
    "浙贝": "浙贝母",
    "丹皮": "牡丹皮",
    "苏叶": "紫苏叶",
    "萎蕤": "玉竹",
    "芍药": "白芍",
}

表达级中药规范 = {
    "制附子钱半": "附子",
    "制附子錢半": "附子",
    "熟附子钱半": "附子",
    "熟附子錢半": "附子",
}

病名规范映射 = {
    "口疮之病": "口疮",
    "吐狂血": "吐血",
    "渗流中道": "",
    "渗留中道": "",
}

剂量单位模式 = re.compile(
    r"^(?P<name>.+?)(?P<dose_unit>(?:\d+(?:\.\d+)?|[零一二三四五六七八九十百千万半两]+)"
    r"(?:mg|mL|ml|g|克|毫克|枚|片|个|条|两|钱|錢|分|厘|升)"
    r"(?:[零一二三四五六七八九十百千万半两]+(?:钱|錢|分|厘))?)(?P<suffix>.*)$"
)
剂量单位拆分模式 = re.compile(
    r"^(?P<dose>\d+(?:\.\d+)?|[零一二三四五六七八九十百千万半两]+)"
    r"(?P<unit>mg|mL|ml|g|克|毫克|枚|片|个|条|两|钱|錢|分|厘|升)"
)
钱半模式 = re.compile(r"^(?P<name>.+?)(?P<dose_unit>钱半|錢半)(?P<suffix>.*)$")
裸单位尾缀模式 = re.compile(r"^(?P<name>.+?)(?P<dose_unit>mg|mL|ml|g|克|毫克)$")
跳过方剂名称 = {"大黄、黄连、菊花、甘草代茶饮"}


def 清理值(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text in 空值集合:
        return ""
    return text


def 清理术语(value: Any) -> str:
    text = 清理值(value)
    text = text.replace("\u3000", " ")
    text = re.sub(r"\s+", "", text)
    return text.strip("“”‘’\"'，,。；;：: ")


def 稳定短哈希(text: str, length: int = 12) -> str:
    return hashlib.sha1(text.encode("utf-8")).hexdigest()[:length]


def 拆分多值字段(value: Any) -> list[str]:
    text = 清理值(value)
    if not text:
        return []
    parts = re.split(r"[、；;\n/]+", text)
    result: list[str] = []
    seen: set[str] = set()
    for part in parts:
        term = 清理术语(part)
        if term and term not in seen:
            result.append(term)
            seen.add(term)
    return result


def 拆分方剂名称字段(value: Any) -> list[str]:
    text = 清理值(value)
    if not text or 清理术语(text) in 跳过方剂名称:
        return []
    parts = re.split(r"[、；;\n/]+", text)
    result: list[str] = []
    seen: set[str] = set()
    for part in parts:
        term = 清理术语(part)
        if term in 跳过方剂名称:
            continue
        if term and term not in seen:
            result.append(term)
            seen.add(term)
    return result


def 读取标准化表(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    records: list[dict[str, str]] = []
    for row in rows[1:]:
        record = {
            header: 清理值(row[index] if index < len(row) else "")
            for index, header in enumerate(headers)
            if header
        }
        if record.get("医案编号"):
            records.append(record)
    return records


def 规范中药名(name: str) -> str:
    name = 清理术语(name)
    if not name:
        return ""
    if name in 中药名规范映射:
        return 中药名规范映射[name]
    name = re.sub(r"生地(?!黄)", "生地黄", name)
    name = re.sub(r"熟地(?!黄)", "熟地黄", name)
    return name


def 规范术语(node_type: str, term: str) -> str:
    term = 清理术语(term)
    if node_type in {"中医病名", "西医病名"}:
        return 病名规范映射.get(term, term)
    return term


def 解析中药表达(raw: str) -> dict[str, str]:
    expression = 清理值(raw)
    clean_expression = 清理术语(expression)
    if clean_expression in 表达级中药规范:
        name = 表达级中药规范[clean_expression]
        return {
            "名称": name,
            "原文术语": clean_expression,
            "剂量": "",
            "单位": "",
            "原文表达": expression,
        }

    match = 钱半模式.match(clean_expression) or 剂量单位模式.match(clean_expression) or 裸单位尾缀模式.match(clean_expression)
    if not match:
        name = clean_expression
        return {
            "名称": 规范中药名(name),
            "原文术语": name,
            "剂量": "",
            "单位": "",
            "原文表达": expression,
        }
    original_name = 清理术语(match.group("name"))
    dose_unit = match.group("dose_unit")
    dose_match = 剂量单位拆分模式.match(dose_unit)
    if dose_unit in {"钱半", "錢半"}:
        dose = ""
        unit = ""
    else:
        dose = ""
        unit = ""
    return {
        "名称": 规范中药名(original_name),
        "原文术语": original_name,
        "剂量": dose,
        "单位": unit,
        "原文表达": expression,
    }


def 节点ID(节点类型: str, 名称: str, 来源医案编号: str = "") -> str:
    if 节点类型 == "医案":
        return f"医案:{名称}"
    if 节点类型 == "患者":
        return f"患者:{来源医案编号}"
    if 节点类型 == "处方":
        return f"处方:{来源医案编号}:1"
    if 节点类型 == "疗效转归":
        return f"疗效转归:{来源医案编号}:{稳定短哈希(名称)}"
    return f"节点:{节点类型}:{稳定短哈希(f'{节点类型}|{名称}')}"


def 关系ID(
    起点节点ID: str,
    关系类型: str,
    终点节点ID: str,
    来源医案编号: str,
    属性: dict[str, str],
) -> str:
    属性串 = "|".join(f"{key}={value}" for key, value in sorted(属性.items()))
    return f"关系:{稳定短哈希(f'{起点节点ID}|{关系类型}|{终点节点ID}|{来源医案编号}|{属性串}', 16)}"


class 图谱构建器:
    def __init__(self) -> None:
        self.nodes: OrderedDict[str, dict[str, str]] = OrderedDict()
        self.relationships: OrderedDict[str, dict[str, str]] = OrderedDict()
        self.triples: list[dict[str, str]] = []

    def add_node(
        self,
        node_type: str,
        name: str,
        *,
        source_case: str = "",
        node_id: str | None = None,
        **attrs: str,
    ) -> str:
        clean_name = 清理值(name)
        if not clean_name:
            raise ValueError(f"禁止创建空名称节点：{node_type}")
        nid = node_id or 节点ID(node_type, clean_name, source_case)
        if nid not in self.nodes:
            self.nodes[nid] = {
                "节点ID": nid,
                "节点类型": node_type,
                "名称": clean_name,
            }
        node = self.nodes[nid]
        for key, value in attrs.items():
            clean = 清理值(value)
            if not clean:
                continue
            if key in {"来源医案编号", "原文术语"} and node.get(key) and clean not in node[key].split("｜"):
                node[key] = f"{node[key]}｜{clean}"
            else:
                node[key] = clean
        return nid

    def add_relationship(
        self,
        start_id: str,
        start_type: str,
        start_name: str,
        rel_type: str,
        end_id: str,
        end_type: str,
        end_name: str,
        source_case: str,
        **attrs: str,
    ) -> None:
        pattern = (start_type, rel_type, end_type)
        allowed_fields = set(
            getattr(kg_schema, "关系属性模式", {}).get(
                pattern,
                getattr(kg_schema, "通用关系属性", []),
            )
        )
        core_fields = {
            "关系ID",
            "起点节点ID",
            "起点节点类型",
            "终点节点ID",
            "终点节点类型",
            "关系类型",
            "来源医案编号",
        }
        allowed_attrs = allowed_fields - core_fields
        rel_attrs = {
            key: 清理值(value)
            for key, value in attrs.items()
            if key in allowed_attrs and 清理值(value)
        }
        rid = 关系ID(start_id, rel_type, end_id, source_case, rel_attrs)
        if rid in self.relationships:
            return
        self.relationships[rid] = {
            "关系ID": rid,
            "起点节点ID": start_id,
            "起点节点类型": start_type,
            "终点节点ID": end_id,
            "终点节点类型": end_type,
            "关系类型": rel_type,
            "来源医案编号": source_case,
            **rel_attrs,
        }
        self.triples.append({
            "医案编号": source_case,
            "头节点ID": start_id,
            "头节点类型": start_type,
            "头实体": start_name,
            "关系": rel_type,
            "尾节点ID": end_id,
            "尾节点类型": end_type,
            "尾实体": end_name,
            **rel_attrs,
        })


def 写CSV(path: Path, fields: list[str], rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def 构建图谱(records: list[dict[str, str]]) -> 图谱构建器:
    graph = 图谱构建器()

    for record in records:
        case_id = record["医案编号"]
        case_name = case_id
        case_node = graph.add_node(
            "医案",
            case_name,
            source_case=case_id,
            医案编号=case_id,
            来源文档=record.get("来源文档", ""),
            既往史=record.get("既往史", ""),
            现病史=record.get("现病史", ""),
        )

        patient_name = 清理值(record.get("姓名")) or f"{case_id}_患者"
        patient_original = "；".join(
            part
            for part in [
                f"姓名={patient_name}",
                f"性别={清理值(record.get('性别')) or 'N'}",
                f"年龄={清理值(record.get('年龄')) or 'N'}",
            ]
        )
        patient_node = graph.add_node(
            "患者",
            patient_name,
            source_case=case_id,
            姓名=patient_name,
            性别=record.get("性别", ""),
            年龄=record.get("年龄", ""),
            来源医案编号=case_id,
            原文表达=patient_original,
        )
        graph.add_relationship(
            case_node,
            "医案",
            case_name,
            "记录患者",
            patient_node,
            "患者",
            patient_name,
            case_id,
            原文表达=patient_original,
        )

        for field_name, node_type, rel_type in 实体字段关系:
            for term in 拆分多值字段(record.get(field_name, "")):
                normalized_term = 规范术语(node_type, term)
                if not normalized_term:
                    continue
                graph_term = normalized_term if node_type in {"中医病名", "西医病名"} else term
                target = graph.add_node(
                    node_type,
                    normalized_term,
                    source_case=case_id,
                    规范名称=normalized_term,
                    原文术语=graph_term,
                    来源医案编号=case_id,
                )
                graph.add_relationship(
                    case_node,
                    "医案",
                    case_name,
                    rel_type,
                    target,
                    node_type,
                    normalized_term,
                    case_id,
                    原文术语=graph_term,
                )

        outcome = 清理值(record.get("疗效转归", ""))
        if outcome:
            outcome_node = graph.add_node(
                "疗效转归",
                outcome,
                source_case=case_id,
                来源医案编号=case_id,
                原文表达=outcome,
            )
            graph.add_relationship(
                case_node,
                "医案",
                case_name,
                "转归为",
                outcome_node,
                "疗效转归",
                outcome,
                case_id,
                原文表达=outcome,
            )

        formula_terms = 拆分方剂名称字段(record.get("方剂名称", ""))
        medicine_terms = 拆分多值字段(record.get("方剂组成", ""))
        formula_node_refs: list[tuple[str, str]] = []
        if formula_terms:
            for formula in formula_terms:
                formula_node = graph.add_node(
                    "方剂",
                    formula,
                    source_case=case_id,
                    规范名称=formula,
                    原文术语=formula,
                )
                graph.add_relationship(
                    case_node,
                    "医案",
                    case_name,
                    "使用方剂",
                    formula_node,
                    "方剂",
                    formula,
                    case_id,
                    原文术语=formula,
                )
                formula_node_refs.append((formula_node, formula))

        if formula_node_refs and medicine_terms:
            formula_node, formula_name = formula_node_refs[0]
            seen_medicines_in_case: set[str] = set()
            for raw_medicine in medicine_terms:
                medicine = 解析中药表达(raw_medicine)
                medicine_name = medicine["名称"]
                if not medicine_name:
                    continue
                dedup_key = f"{formula_node}|{medicine_name}"
                if dedup_key in seen_medicines_in_case:
                    continue
                seen_medicines_in_case.add(dedup_key)
                medicine_node = graph.add_node(
                    "中药",
                    medicine_name,
                    source_case=case_id,
                    规范名称=medicine_name,
                    原文术语=medicine["原文术语"],
                )
                graph.add_relationship(
                    formula_node,
                    "方剂",
                    formula_name,
                    "包含中药",
                    medicine_node,
                    "中药",
                    medicine_name,
                    case_id,
                    原文术语=medicine["原文术语"],
                    剂量=medicine["剂量"],
                    单位=medicine["单位"],
                    原文表达=medicine["原文表达"],
                )

    return graph


def schema_有效节点类型() -> set[str]:
    relation_types = {head for head, _, _ in kg_schema.关系模式白名单} | {
        tail for _, _, tail in kg_schema.关系模式白名单
    }
    property_types = set(kg_schema.节点属性模式)
    return set(kg_schema.节点类型白名单) | relation_types | property_types


def 校验schema(graph: 图谱构建器) -> list[str]:
    warnings: list[str] = []
    raw_node_types = set(kg_schema.节点类型白名单)
    effective_node_types = schema_有效节点类型()
    schema_internal_missing = sorted(effective_node_types - raw_node_types)
    if schema_internal_missing:
        warnings.append(
            "kg_schema.py 节点类型白名单未列出但关系/属性模式引用："
            + "、".join(schema_internal_missing)
        )

    bad_node_types = sorted({
        row["节点类型"]
        for row in graph.nodes.values()
        if row["节点类型"] not in effective_node_types
    })
    if bad_node_types:
        raise ValueError(f"发现 schema 外节点类型：{bad_node_types}")

    schema_relations = set(kg_schema.关系模式白名单)
    bad_relations = sorted({
        (row["起点节点类型"], row["关系类型"], row["终点节点类型"])
        for row in graph.relationships.values()
        if (row["起点节点类型"], row["关系类型"], row["终点节点类型"]) not in schema_relations
    })
    if bad_relations:
        raise ValueError(f"发现 schema 外关系模式：{bad_relations[:20]}")

    relation_attr_schema = getattr(kg_schema, "关系属性模式", {})
    bad_relation_attrs = []
    for row in graph.relationships.values():
        pattern = (row["起点节点类型"], row["关系类型"], row["终点节点类型"])
        allowed = set(relation_attr_schema.get(pattern, getattr(kg_schema, "通用关系属性", [])))
        extra = sorted(key for key, value in row.items() if 清理值(value) and key not in allowed)
        if extra:
            bad_relation_attrs.append((row["关系ID"], pattern, extra))
    if bad_relation_attrs:
        raise ValueError(f"发现 schema 外关系属性：{bad_relation_attrs[:20]}")

    node_ids = set(graph.nodes)
    missing_endpoint = [
        row["关系ID"]
        for row in graph.relationships.values()
        if row["起点节点ID"] not in node_ids or row["终点节点ID"] not in node_ids
    ]
    if missing_endpoint:
        raise ValueError(f"发现缺失端点关系：{missing_endpoint[:20]}")

    return warnings


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="按 kg_schema.py 构建中文图谱 CSV")
    parser.add_argument("--input", type=Path, required=True, help="人工确认后的标准化工作簿")
    parser.add_argument("--mapping", type=Path, default=默认映射路径)
    parser.add_argument("--output-dir", type=Path, default=默认输出目录)
    parser.add_argument("--date", default=datetime.now().strftime("%Y%m%d_%H%M%S"))
    parser.add_argument("--suffix", default="graph")
    return parser.parse_args()


def 校验医案来源(records: list[dict[str, str]], mapping_path: Path) -> None:
    """拒绝重复编号和未经映射确认的来源；允许只处理可用的子集。"""
    if not mapping_path.is_file():
        raise FileNotFoundError(f"缺少阶段02来源映射：{mapping_path}")
    with mapping_path.open("r", encoding="utf-8-sig", newline="") as source:
        rows = list(csv.DictReader(source))
    mapping = {row["医案编号"]: row["文档名称"] for row in rows}
    if len(mapping) != len(rows):
        raise ValueError("来源映射存在重复 Case ID")
    seen: set[str] = set()
    for record in records:
        case_id = record["医案编号"]
        if not re.fullmatch(r"Case\d+", case_id) or case_id in seen:
            raise ValueError(f"无效或重复的医案编号：{case_id}")
        seen.add(case_id)
        if case_id not in mapping or record.get("来源文档", "") != mapping[case_id]:
            raise ValueError(f"医案来源与阶段02映射不一致：{case_id}")


def main() -> None:
    args = parse_args()
    records = 读取标准化表(args.input)
    if not records:
        raise FileNotFoundError(f"没有读到标准化表记录：{args.input}")
    校验医案来源(records, args.mapping)

    graph = 构建图谱(records)
    warnings = 校验schema(graph)

    output_dir = args.output_dir / f"{args.date}_{args.suffix}"
    output_dir.mkdir(parents=True, exist_ok=False)
    nodes_path = output_dir / "节点表.csv"
    relationships_path = output_dir / "关系表.csv"
    triples_path = output_dir / "三元组表.csv"

    node_fields = list(getattr(kg_schema, "节点表字段", []))
    if not node_fields:
        for fields in kg_schema.节点属性模式.values():
            for field in fields:
                if field not in node_fields:
                    node_fields.append(field)

    relationship_fields = list(getattr(kg_schema, "关系表字段", []))
    if not relationship_fields:
        relationship_fields = [
            "关系ID",
            "起点节点ID",
            "起点节点类型",
            "终点节点ID",
            "终点节点类型",
            "关系类型",
            "来源医案编号",
            "原文术语",
            "原文表达",
            "剂量",
            "单位",
        ]
    triple_fields = [
        "医案编号",
        "头节点ID",
        "头节点类型",
        "头实体",
        "关系",
        "尾节点ID",
        "尾节点类型",
        "尾实体",
        "原文术语",
        "原文表达",
        "剂量",
        "单位",
    ]

    node_rows = sorted(
        graph.nodes.values(),
        key=lambda row: (row.get("节点类型", ""), row.get("名称", ""), row.get("节点ID", "")),
    )
    relationship_rows = sorted(
        graph.relationships.values(),
        key=lambda row: (row.get("来源医案编号", ""), row.get("关系类型", ""), row.get("关系ID", "")),
    )

    写CSV(nodes_path, node_fields, node_rows)
    写CSV(relationships_path, relationship_fields, relationship_rows)
    写CSV(triples_path, triple_fields, graph.triples)

    counts_by_node_type: dict[str, int] = {}
    for row in node_rows:
        counts_by_node_type[row["节点类型"]] = counts_by_node_type.get(row["节点类型"], 0) + 1

    counts_by_relation: dict[str, int] = {}
    for row in relationship_rows:
        counts_by_relation[row["关系类型"]] = counts_by_relation.get(row["关系类型"], 0) + 1

    print("=" * 72)
    print("kg_schema 新版 CSV 生成完成")
    print(f"输入标准化表：{args.input}")
    print(f"医案记录数：{len(records)}")
    print(f"节点数：{len(node_rows)}")
    print(f"关系数：{len(relationship_rows)}")
    print(f"三元组数：{len(graph.triples)}")
    print(f"患者节点数：{counts_by_node_type.get('患者', 0)}")
    print(f"记录患者关系数：{counts_by_relation.get('记录患者', 0)}")
    for warning in warnings:
        print(f"[schema提示] {warning}")
    print(f"nodes：{nodes_path}")
    print(f"relationships：{relationships_path}")
    print(f"triples：{triples_path}")
    print("=" * 72)


if __name__ == "__main__":
    main()
